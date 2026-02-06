from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List

from openpyxl import Workbook, load_workbook

from .data_sources import SensorSample


@dataclass
class RecordingPaths:
    folder: Path
    active_file: Path


def _ensure_workbook(path: Path, sheet_names: Iterable[str]) -> None:
    if path.exists():
        return
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    for name in sheet_names:
        sheet = workbook.create_sheet(title=name)
        sheet.append(["Tag name", "Time (hh:mm:ss)", "Value"])
    workbook.save(path)


def create_recording_paths(base_path: Path, timestamp: datetime) -> RecordingPaths:
    folder = base_path / timestamp.strftime("%Y-%m-%d")
    folder.mkdir(parents=True, exist_ok=True)
    active_file = folder / "recording.xlsx"
    return RecordingPaths(folder=folder, active_file=active_file)


def append_rows(path: Path, rows_by_sheet: Dict[str, List[SensorSample]]) -> None:
    workbook = load_workbook(path)
    for sheet_name, samples in rows_by_sheet.items():
        sheet = workbook[sheet_name]
        for sample in samples:
            sheet.append(
                [
                    sample.tag,
                    sample.timestamp.strftime("%H:%M:%S"),
                    sample.value,
                ]
            )
    workbook.save(path)


def finalize_recording(active_file: Path, start_time: datetime, end_time: datetime) -> Path:
    final_name = f"{start_time.strftime('%H%M%S')}_{end_time.strftime('%H%M%S')}.xlsx"
    final_path = active_file.with_name(final_name)
    active_file.rename(final_path)
    return final_path
