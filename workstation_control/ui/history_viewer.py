"""
History viewer dialog for loading and visualizing previously recorded Excel files.

Allows users to:
- Browse date folders and select recorded Excel files
- Display the data from Temperature, Pressure, Flow sheets
- Plot historical trends using pyqtgraph
"""

import os

import pyqtgraph as pg
from openpyxl import load_workbook
from PyQt5.QtCore import Qt
from PyQt5.QtWidgets import (
    QComboBox,
    QDialog,
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QPushButton,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
)

COLORS = ["#2196F3", "#F44336", "#4CAF50", "#FF9800", "#9C27B0", "#00BCD4"]


class HistoryViewer(QDialog):
    """Dialog for viewing historical recorded data."""

    def __init__(self, base_dir: str = ".", parent=None):
        super().__init__(parent)
        self._base_dir = base_dir
        self.setWindowTitle("History Viewer")
        self.setMinimumSize(900, 600)
        self._build_ui()

    def _build_ui(self):
        layout = QVBoxLayout(self)

        # File selection toolbar
        toolbar = QHBoxLayout()
        self._open_btn = QPushButton("Open File...")
        self._open_btn.clicked.connect(self._open_file)
        toolbar.addWidget(self._open_btn)

        self._file_label = QLabel("No file loaded")
        self._file_label.setStyleSheet("color: #666;")
        toolbar.addWidget(self._file_label)
        toolbar.addStretch()

        # Sheet selector
        toolbar.addWidget(QLabel("Sheet:"))
        self._sheet_combo = QComboBox()
        self._sheet_combo.addItems(["Temperature", "Pressure", "Flow"])
        self._sheet_combo.currentTextChanged.connect(self._on_sheet_changed)
        toolbar.addWidget(self._sheet_combo)

        layout.addLayout(toolbar)

        # Main content: splitter with table and chart
        splitter = QSplitter(Qt.Vertical)

        # Chart
        pg.setConfigOptions(antialias=True, background="w", foreground="k")
        self._plot = pg.PlotWidget(title="Historical Data")
        self._plot.showGrid(x=True, y=True, alpha=0.3)
        self._plot.setLabel("bottom", "Sample Index")
        self._plot.addLegend()
        splitter.addWidget(self._plot)

        # Table
        self._table = QTableWidget()
        self._table.setColumnCount(4)
        self._table.setHorizontalHeaderLabels(["Tag", "Time", "Value", "Unit"])
        self._table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self._table.setAlternatingRowColors(True)
        splitter.addWidget(self._table)

        splitter.setSizes([400, 200])
        layout.addWidget(splitter)

        # Internal data cache
        self._sheets_data = {}

    def _open_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Open Recording",
            self._base_dir,
            "Excel Files (*.xlsx)",
        )
        if not path:
            return

        self._load_file(path)

    def _load_file(self, path: str):
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            self._sheets_data.clear()

            for sheet_name in ["Temperature", "Pressure", "Flow"]:
                if sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    rows = []
                    for i, row in enumerate(ws.iter_rows(values_only=True)):
                        if i == 0:
                            continue  # skip header
                        rows.append(row)
                    self._sheets_data[sheet_name] = rows

            wb.close()
            self._file_label.setText(os.path.basename(path))
            self._on_sheet_changed(self._sheet_combo.currentText())

        except Exception as e:
            self._file_label.setText(f"Error: {e}")

    def _on_sheet_changed(self, sheet_name: str):
        rows = self._sheets_data.get(sheet_name, [])

        # Update table
        self._table.setRowCount(len(rows))
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                item = QTableWidgetItem(str(val) if val is not None else "")
                self._table.setItem(i, j, item)

        # Update chart
        self._plot.clear()
        if not rows:
            return

        # Group data by tag
        tag_data = {}
        for row in rows:
            tag = row[0] if row[0] else "Unknown"
            val = row[2] if row[2] is not None else 0
            try:
                val = float(val)
            except (ValueError, TypeError):
                val = 0
            tag_data.setdefault(tag, []).append(val)

        for idx, (tag, values) in enumerate(tag_data.items()):
            color = COLORS[idx % len(COLORS)]
            pen = pg.mkPen(color=color, width=2)
            self._plot.plot(
                list(range(len(values))),
                values,
                pen=pen,
                name=tag,
            )

        self._plot.setTitle(f"{sheet_name} - Historical Data")
