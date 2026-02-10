#!/usr/bin/env python3
"""Unit tests for all workstation_control modules."""

import json
import os
import sys
import tempfile
import shutil
import time
from datetime import datetime

# Add project root to path
sys.path.insert(0, os.path.dirname(__file__))

# ============================================================
# Test 1: config_loader
# ============================================================
def test_config_loader():
    print("=" * 60)
    print("TEST 1: config_loader")
    print("=" * 60)
    from core.config_loader import load_config, ConfigError

    errors = []

    # 1a: Load valid config
    print("\n[1a] Load valid config (station_2pumps.json)...")
    try:
        cfg = load_config("config/station_2pumps.json")
        assert cfg["station_name"] == "Workstation A - 2 Pumps"
        assert cfg["plc"]["host"] == "192.168.1.10"
        assert cfg["plc"]["port"] == 502
        assert len(cfg["devices"]) == 3
        assert len(cfg["sensors"]) == 5
        print("  PASS: Valid config loaded successfully")
    except Exception as e:
        errors.append(f"1a: {e}")
        print(f"  FAIL: {e}")

    # 1b: Load second valid config
    print("\n[1b] Load valid config (station_3pumps.json)...")
    try:
        cfg = load_config("config/station_3pumps.json")
        assert cfg["station_name"] == "Workstation B - 3 Pumps"
        assert len(cfg["devices"]) == 5
        assert len(cfg["sensors"]) == 8
        print("  PASS: Valid config loaded successfully")
    except Exception as e:
        errors.append(f"1b: {e}")
        print(f"  FAIL: {e}")

    # 1c: Defaults applied
    print("\n[1c] Check default values applied...")
    try:
        cfg = load_config("config/station_2pumps.json")
        assert cfg["poll_interval_ms"] == 1000
        for sensor in cfg["sensors"]:
            assert "scale" in sensor
            assert "offset" in sensor
            assert "data_type" in sensor
            assert "alarm_high" in sensor
            assert "alarm_low" in sensor
        for device in cfg["devices"]:
            assert "setpoint_register" in device
            assert "setpoint_min" in device
            assert "setpoint_max" in device
        print("  PASS: All defaults applied correctly")
    except Exception as e:
        errors.append(f"1c: {e}")
        print(f"  FAIL: {e}")

    # 1d: Missing file
    print("\n[1d] Missing config file...")
    try:
        load_config("nonexistent.json")
        errors.append("1d: Should have raised ConfigError")
        print("  FAIL: No exception raised")
    except ConfigError as e:
        print(f"  PASS: ConfigError raised: {e}")
    except Exception as e:
        errors.append(f"1d: Wrong exception type: {e}")
        print(f"  FAIL: Wrong exception: {e}")

    # 1e: Invalid config - missing station_name
    print("\n[1e] Invalid config - missing station_name...")
    with tempfile.NamedTemporaryFile(mode='w', suffix='.json', delete=False) as f:
        json.dump({"plc": {"host": "1.1.1.1", "port": 502, "unit_id": 1}}, f)
        tmp_path = f.name
    try:
        load_config(tmp_path)
        errors.append("1e: Should have raised ConfigError")
        print("  FAIL: No exception raised")
    except ConfigError as e:
        print(f"  PASS: ConfigError raised: {e}")
    except Exception as e:
        errors.append(f"1e: Wrong exception: {e}")
        print(f"  FAIL: {e}")
    finally:
        os.unlink(tmp_path)

    # 1f: Invalid config - missing devices
    print("\n[1f] Invalid config - no devices...")
    with tempfile.NamedTemporaryFile(mode='w', suffix='.json', delete=False) as f:
        json.dump({
            "station_name": "Test",
            "plc": {"host": "1.1.1.1", "port": 502, "unit_id": 1},
            "devices": [],
            "sensors": [{"tag": "T1", "name": "T", "category": "Temperature", "register": 1, "unit": "C"}]
        }, f)
        tmp_path = f.name
    try:
        load_config(tmp_path)
        errors.append("1f: Should have raised ConfigError")
        print("  FAIL: No exception raised")
    except ConfigError as e:
        print(f"  PASS: ConfigError raised: {e}")
    finally:
        os.unlink(tmp_path)

    # 1g: Invalid device type
    print("\n[1g] Invalid device type...")
    with tempfile.NamedTemporaryFile(mode='w', suffix='.json', delete=False) as f:
        json.dump({
            "station_name": "Test",
            "plc": {"host": "1.1.1.1", "port": 502, "unit_id": 1},
            "devices": [{"name": "X", "type": "robot", "control_register": 0, "status_register": 1}],
            "sensors": [{"tag": "T1", "name": "T", "category": "Temperature", "register": 1, "unit": "C"}]
        }, f)
        tmp_path = f.name
    try:
        load_config(tmp_path)
        errors.append("1g: Should have raised ConfigError")
        print("  FAIL: No exception raised")
    except ConfigError as e:
        print(f"  PASS: ConfigError raised: {e}")
    finally:
        os.unlink(tmp_path)

    # 1h: Invalid sensor category
    print("\n[1h] Invalid sensor category...")
    with tempfile.NamedTemporaryFile(mode='w', suffix='.json', delete=False) as f:
        json.dump({
            "station_name": "Test",
            "plc": {"host": "1.1.1.1", "port": 502, "unit_id": 1},
            "devices": [{"name": "P1", "type": "pump", "control_register": 0, "status_register": 1}],
            "sensors": [{"tag": "T1", "name": "T", "category": "Humidity", "register": 1, "unit": "%"}]
        }, f)
        tmp_path = f.name
    try:
        load_config(tmp_path)
        errors.append("1h: Should have raised ConfigError")
        print("  FAIL: No exception raised")
    except ConfigError as e:
        print(f"  PASS: ConfigError raised: {e}")
    finally:
        os.unlink(tmp_path)

    return errors


# ============================================================
# Test 2: data_recorder
# ============================================================
def test_data_recorder():
    print("\n" + "=" * 60)
    print("TEST 2: data_recorder")
    print("=" * 60)

    errors = []
    tmp_dir = tempfile.mkdtemp()

    try:
        # Need QApplication for signals
        from PyQt5.QtWidgets import QApplication
        app = QApplication.instance()
        if app is None:
            app = QApplication(sys.argv)

        from core.data_recorder import DataRecorder
        from core.modbus_client import SensorReading

        # 2a: Date folder creation
        print("\n[2a] Date folder creation on init...")
        recorder = DataRecorder(base_dir=tmp_dir)
        today = datetime.now().strftime("%Y-%m-%d")
        expected_folder = os.path.join(tmp_dir, today)
        assert os.path.isdir(expected_folder), f"Folder not created: {expected_folder}"
        print(f"  PASS: Date folder created: {expected_folder}")

        # 2b: Start recording
        print("\n[2b] Start recording...")
        result = recorder.start_recording()
        assert result is True, "start_recording returned False"
        assert recorder.is_recording is True
        temp_file = os.path.join(expected_folder, "_recording_in_progress.xlsx")
        assert os.path.isfile(temp_file), f"Temp file not created: {temp_file}"
        print(f"  PASS: Recording started, temp file: {temp_file}")

        # 2c: Verify Excel sheets
        print("\n[2c] Verify Excel has 3 sheets...")
        from openpyxl import load_workbook
        wb = load_workbook(temp_file)
        sheet_names = wb.sheetnames
        wb.close()
        assert "Temperature" in sheet_names, f"Missing Temperature sheet: {sheet_names}"
        assert "Pressure" in sheet_names, f"Missing Pressure sheet: {sheet_names}"
        assert "Flow" in sheet_names, f"Missing Flow sheet: {sheet_names}"
        print(f"  PASS: Sheets found: {sheet_names}")

        # 2d: Append data
        print("\n[2d] Append sensor data...")
        now_str = datetime.now().strftime("%H:%M:%S")
        test_readings = [
            SensorReading("TC-001", "Temp1", "Temperature", 45.5, "°C", now_str),
            SensorReading("PT-001", "Pres1", "Pressure", 3.2, "bar", now_str),
            SensorReading("FT-001", "Flow1", "Flow", 12.1, "L/min", now_str),
        ]
        recorder.append_data(test_readings)

        # Append a few more rounds
        for _ in range(5):
            recorder.append_data(test_readings)

        print("  PASS: Data appended without errors")

        # 2e: Prevent double start
        print("\n[2e] Prevent double start...")
        result2 = recorder.start_recording()
        assert result2 is False, "Double start should return False"
        print("  PASS: Double start correctly rejected")

        # 2f: Stop recording and verify file naming
        print("\n[2f] Stop recording...")
        final_path = recorder.stop_recording()
        assert final_path is not None, "stop_recording returned None"
        assert recorder.is_recording is False
        assert os.path.isfile(final_path), f"Final file not found: {final_path}"
        basename = os.path.basename(final_path)
        assert basename.endswith(".xlsx"), f"Wrong extension: {basename}"
        # Check format: HHMMSS_HHMMSS.xlsx
        name_part = basename.replace(".xlsx", "")
        parts = name_part.split("_")
        assert len(parts) >= 2, f"Unexpected filename format: {basename}"
        assert len(parts[0]) == 6, f"Start time wrong format: {parts[0]}"
        assert len(parts[1]) == 6, f"End time wrong format: {parts[1]}"
        print(f"  PASS: File renamed to: {basename}")

        # 2g: Verify data in final file
        print("\n[2g] Verify data content in saved file...")
        wb = load_workbook(final_path)
        temp_ws = wb["Temperature"]
        rows = list(temp_ws.iter_rows(values_only=True))
        assert rows[0] == ("Tag", "Time", "Value", "Unit"), f"Header mismatch: {rows[0]}"
        assert len(rows) > 1, "No data rows found"
        assert rows[1][0] == "TC-001", f"Tag mismatch: {rows[1][0]}"

        pres_ws = wb["Pressure"]
        pres_rows = list(pres_ws.iter_rows(values_only=True))
        assert len(pres_rows) > 1
        assert pres_rows[1][0] == "PT-001"

        flow_ws = wb["Flow"]
        flow_rows = list(flow_ws.iter_rows(values_only=True))
        assert len(flow_rows) > 1
        assert flow_rows[1][0] == "FT-001"

        wb.close()
        print(f"  PASS: Temperature={len(rows)-1} rows, Pressure={len(pres_rows)-1} rows, Flow={len(flow_rows)-1} rows")

        # 2h: Stop when not recording
        print("\n[2h] Stop when not recording...")
        result3 = recorder.stop_recording()
        assert result3 is None, "Should return None when not recording"
        print("  PASS: Correctly returns None")

    except Exception as e:
        errors.append(f"2: {e}")
        import traceback
        traceback.print_exc()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    return errors


# ============================================================
# Test 3: modbus_client (decode / alarm logic)
# ============================================================
def test_modbus_client_logic():
    print("\n" + "=" * 60)
    print("TEST 3: modbus_client (decode & alarm logic)")
    print("=" * 60)

    errors = []

    try:
        from PyQt5.QtWidgets import QApplication
        app = QApplication.instance()
        if app is None:
            app = QApplication(sys.argv)

        from core.modbus_client import ModbusClient

        cfg = {
            "station_name": "Test",
            "plc": {"host": "127.0.0.1", "port": 502, "unit_id": 1},
            "poll_interval_ms": 1000,
            "devices": [{"name": "P1", "type": "pump", "control_register": 0,
                         "status_register": 100, "setpoint_register": 200,
                         "setpoint_unit": "RPM", "setpoint_min": 0, "setpoint_max": 3000}],
            "sensors": [{"tag": "TC-001", "name": "T1", "category": "Temperature",
                         "register": 300, "data_type": "uint16", "scale": 0.1,
                         "offset": 0, "unit": "°C", "alarm_high": 100.0, "alarm_low": 5.0}],
        }

        client = ModbusClient(cfg)

        # 3a: _decode_registers - uint16
        print("\n[3a] Decode uint16...")
        val = client._decode_registers([500], "uint16")
        assert val == 500, f"Expected 500, got {val}"
        print(f"  PASS: uint16 decode = {val}")

        # 3b: _decode_registers - int16 positive
        print("\n[3b] Decode int16 (positive)...")
        val = client._decode_registers([100], "int16")
        assert val == 100, f"Expected 100, got {val}"
        print(f"  PASS: int16 positive = {val}")

        # 3c: _decode_registers - int16 negative
        print("\n[3c] Decode int16 (negative)...")
        val = client._decode_registers([65535], "int16")
        assert val == -1, f"Expected -1, got {val}"
        val2 = client._decode_registers([65036], "int16")
        assert val2 == -500, f"Expected -500, got {val2}"
        print(f"  PASS: int16 negative: 65535->{val}, 65036->{val2}")

        # 3d: _decode_registers - int32
        print("\n[3d] Decode int32...")
        # 100000 in big-endian int32: 0x000186A0 -> registers [0x0001, 0x86A0]
        val = client._decode_registers([0x0001, 0x86A0], "int32")
        assert val == 100000, f"Expected 100000, got {val}"
        print(f"  PASS: int32 = {val}")

        # 3e: _decode_registers - float32
        print("\n[3e] Decode float32...")
        import struct
        raw = struct.pack(">f", 3.14)
        regs = struct.unpack(">HH", raw)
        val = client._decode_registers(list(regs), "float32")
        assert abs(val - 3.14) < 0.001, f"Expected ~3.14, got {val}"
        print(f"  PASS: float32 = {val:.4f}")

        # 3f: _check_alarm - HIGH
        print("\n[3f] Alarm check - HIGH...")
        sensor = {"alarm_high": 100.0, "alarm_low": 5.0}
        result = client._check_alarm(sensor, 105.0)
        assert result == "HIGH", f"Expected HIGH, got {result}"
        print(f"  PASS: alarm={result}")

        # 3g: _check_alarm - LOW
        print("\n[3g] Alarm check - LOW...")
        result = client._check_alarm(sensor, 3.0)
        assert result == "LOW", f"Expected LOW, got {result}"
        print(f"  PASS: alarm={result}")

        # 3h: _check_alarm - None (normal)
        print("\n[3h] Alarm check - Normal...")
        result = client._check_alarm(sensor, 50.0)
        assert result is None, f"Expected None, got {result}"
        print(f"  PASS: alarm={result}")

        # 3i: _check_alarm - no thresholds
        print("\n[3i] Alarm check - No thresholds...")
        result = client._check_alarm({}, 50.0)
        assert result is None
        print(f"  PASS: alarm={result}")

    except Exception as e:
        errors.append(f"3: {e}")
        import traceback
        traceback.print_exc()

    return errors


# ============================================================
# Test 4: trend_chart widget
# ============================================================
def test_trend_chart():
    print("\n" + "=" * 60)
    print("TEST 4: trend_chart widget")
    print("=" * 60)

    errors = []

    try:
        from PyQt5.QtWidgets import QApplication
        app = QApplication.instance()
        if app is None:
            app = QApplication(sys.argv)

        from ui.trend_chart import TrendChart
        from core.modbus_client import SensorReading

        sensors = [
            {"tag": "TC-001", "name": "Temp1", "category": "Temperature", "unit": "°C",
             "register": 300, "data_type": "uint16", "scale": 0.1, "offset": 0},
            {"tag": "PT-001", "name": "Pres1", "category": "Pressure", "unit": "bar",
             "register": 301, "data_type": "uint16", "scale": 0.01, "offset": 0},
            {"tag": "FT-001", "name": "Flow1", "category": "Flow", "unit": "L/min",
             "register": 302, "data_type": "uint16", "scale": 0.1, "offset": 0},
        ]

        # 4a: Widget creation
        print("\n[4a] Create TrendChart widget...")
        chart = TrendChart(sensors)
        assert chart is not None
        assert len(chart._plots) == 3
        assert "Temperature" in chart._plots
        assert "Pressure" in chart._plots
        assert "Flow" in chart._plots
        assert len(chart._curves) == 3
        print(f"  PASS: 3 plots, 3 curves created")

        # 4b: Feed data
        print("\n[4b] Feed data points...")
        for i in range(10):
            readings = [
                SensorReading("TC-001", "Temp1", "Temperature", 45.0 + i * 0.1, "°C", "12:00:00"),
                SensorReading("PT-001", "Pres1", "Pressure", 3.0 + i * 0.01, "bar", "12:00:00"),
                SensorReading("FT-001", "Flow1", "Flow", 12.0 + i * 0.1, "L/min", "12:00:00"),
            ]
            chart.update_data(readings)

        assert chart._tick_counter == 10
        assert len(chart._data["TC-001"]) == 10
        assert len(chart._data["PT-001"]) == 10
        assert len(chart._data["FT-001"]) == 10
        print(f"  PASS: 10 data points per sensor, tick={chart._tick_counter}")

        # 4c: Clear data
        print("\n[4c] Clear data...")
        chart.clear_data()
        assert chart._tick_counter == 0
        assert len(chart._data) == 0
        print("  PASS: Data cleared")

        # 4d: Window change
        print("\n[4d] Window time change...")
        chart._on_window_changed("1 min")
        assert chart._window_seconds == 60
        chart._on_window_changed("30 min")
        assert chart._window_seconds == 1800
        print(f"  PASS: Window changes work correctly")

    except Exception as e:
        errors.append(f"4: {e}")
        import traceback
        traceback.print_exc()

    return errors


# ============================================================
# Test 5: control_panel widget
# ============================================================
def test_control_panel():
    print("\n" + "=" * 60)
    print("TEST 5: control_panel widget")
    print("=" * 60)

    errors = []

    try:
        from PyQt5.QtWidgets import QApplication
        app = QApplication.instance()
        if app is None:
            app = QApplication(sys.argv)

        from ui.control_panel import ControlPanel, DeviceCard
        from core.modbus_client import SensorReading

        devices = [
            {"name": "Pump 1", "type": "pump", "control_register": 0,
             "status_register": 100, "setpoint_register": 200,
             "setpoint_unit": "RPM", "setpoint_min": 0, "setpoint_max": 3000},
            {"name": "Heater 1", "type": "heater", "control_register": 2,
             "status_register": 102, "setpoint_register": 202,
             "setpoint_unit": "°C", "setpoint_min": 0, "setpoint_max": 200},
        ]
        sensors = [
            {"tag": "TC-001", "name": "Temp1", "category": "Temperature", "unit": "°C",
             "register": 300, "data_type": "uint16", "scale": 0.1, "offset": 0},
            {"tag": "PT-001", "name": "Pres1", "category": "Pressure", "unit": "bar",
             "register": 301, "data_type": "uint16", "scale": 0.01, "offset": 0},
        ]

        # 5a: Panel creation
        print("\n[5a] Create ControlPanel...")
        panel = ControlPanel(devices, sensors)
        assert len(panel._cards) == 2
        assert "Pump 1" in panel._cards
        assert "Heater 1" in panel._cards
        print(f"  PASS: 2 device cards created")

        # 5b: Device status update
        print("\n[5b] Update device status...")
        panel.update_device_status("Pump 1", True)
        card = panel._cards["Pump 1"]
        assert card._is_on is True
        assert card._toggle_btn.text() == "STOP"
        assert card._status_label.text() == "ON"
        print("  PASS: Pump 1 status=ON, button=STOP")

        panel.update_device_status("Pump 1", False)
        assert card._is_on is False
        assert card._toggle_btn.text() == "START"
        assert card._status_label.text() == "OFF"
        print("  PASS: Pump 1 status=OFF, button=START")

        # 5c: Sensor readings update
        print("\n[5c] Update sensor readings...")
        readings = [
            SensorReading("TC-001", "Temp1", "Temperature", 78.55, "°C", "12:00:00"),
            SensorReading("PT-001", "Pres1", "Pressure", 5.12, "bar", "12:00:00", alarm="HIGH"),
        ]
        panel.update_sensor_readings(readings)
        tc_info = panel._sensor_labels["TC-001"]
        assert tc_info["value"].text() == "78.55"
        pt_info = panel._sensor_labels["PT-001"]
        assert pt_info["value"].text() == "5.12"
        assert pt_info["status"].text() == "HIGH"
        print("  PASS: Sensor values and alarm status updated")

        # 5d: DeviceCard setpoint spin box range
        print("\n[5d] Setpoint spin box range check...")
        pump_card = panel._cards["Pump 1"]
        assert pump_card._setpoint_spin is not None
        assert pump_card._setpoint_spin.minimum() == 0
        assert pump_card._setpoint_spin.maximum() == 3000
        heater_card = panel._cards["Heater 1"]
        assert heater_card._setpoint_spin.maximum() == 200
        print("  PASS: Spin box ranges correct")

        # 5e: Signal emission
        print("\n[5e] Signal emission test...")
        received = []
        panel.device_toggle.connect(lambda dev, state: received.append(("toggle", dev["name"], state)))
        panel.device_setpoint.connect(lambda dev, val: received.append(("setpoint", dev["name"], val)))

        pump_card._setpoint_spin.setValue(1500.0)
        pump_card._set_btn.click()
        assert len(received) == 1
        assert received[0] == ("setpoint", "Pump 1", 1500.0)
        print(f"  PASS: Setpoint signal emitted: {received[0]}")

    except Exception as e:
        errors.append(f"5: {e}")
        import traceback
        traceback.print_exc()

    return errors


# ============================================================
# Test 6: history_viewer
# ============================================================
def test_history_viewer():
    print("\n" + "=" * 60)
    print("TEST 6: history_viewer")
    print("=" * 60)

    errors = []
    tmp_dir = tempfile.mkdtemp()

    try:
        from PyQt5.QtWidgets import QApplication
        app = QApplication.instance()
        if app is None:
            app = QApplication(sys.argv)

        from ui.history_viewer import HistoryViewer
        from openpyxl import Workbook

        # Create a test Excel file
        test_file = os.path.join(tmp_dir, "120000_121000.xlsx")
        wb = Workbook()
        default = wb.active
        wb.remove(default)

        for sheet_name, tag, values in [
            ("Temperature", "TC-001", [45.1, 45.2, 45.3, 45.5, 45.8]),
            ("Pressure", "PT-001", [3.1, 3.2, 3.3, 3.2, 3.1]),
            ("Flow", "FT-001", [12.0, 12.1, 12.2, 12.3, 12.4]),
        ]:
            ws = wb.create_sheet(title=sheet_name)
            ws.append(["Tag", "Time", "Value", "Unit"])
            for i, val in enumerate(values):
                ws.append([tag, f"12:00:0{i}", val, "unit"])

        wb.save(test_file)
        wb.close()

        # 6a: Create viewer
        print("\n[6a] Create HistoryViewer...")
        viewer = HistoryViewer(base_dir=tmp_dir)
        assert viewer is not None
        print("  PASS: HistoryViewer created")

        # 6b: Load file programmatically
        print("\n[6b] Load test file...")
        viewer._load_file(test_file)
        assert "Temperature" in viewer._sheets_data
        assert "Pressure" in viewer._sheets_data
        assert "Flow" in viewer._sheets_data
        assert len(viewer._sheets_data["Temperature"]) == 5
        assert len(viewer._sheets_data["Pressure"]) == 5
        assert len(viewer._sheets_data["Flow"]) == 5
        print(f"  PASS: Loaded 5 rows per sheet")

        # 6c: Verify table population
        print("\n[6c] Switch to Temperature sheet...")
        viewer._on_sheet_changed("Temperature")
        assert viewer._table.rowCount() == 5
        assert viewer._table.item(0, 0).text() == "TC-001"
        print(f"  PASS: Table shows 5 rows, first tag=TC-001")

        # 6d: Switch sheets
        print("\n[6d] Switch to Pressure sheet...")
        viewer._on_sheet_changed("Pressure")
        assert viewer._table.rowCount() == 5
        assert viewer._table.item(0, 0).text() == "PT-001"
        print(f"  PASS: Pressure sheet shows PT-001")

        # 6e: File label
        print("\n[6e] File label updated...")
        assert viewer._file_label.text() == "120000_121000.xlsx"
        print(f"  PASS: Label = {viewer._file_label.text()}")

    except Exception as e:
        errors.append(f"6: {e}")
        import traceback
        traceback.print_exc()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    return errors


# ============================================================
# Test 7: Simulation mode integration
# ============================================================
def test_simulation_integration():
    print("\n" + "=" * 60)
    print("TEST 7: Simulation mode integration")
    print("=" * 60)

    errors = []
    tmp_dir = tempfile.mkdtemp()

    try:
        from PyQt5.QtWidgets import QApplication
        from PyQt5.QtCore import QTimer
        app = QApplication.instance()
        if app is None:
            app = QApplication(sys.argv)

        from core.config_loader import load_config

        # 7a: Load config and patch simulation
        print("\n[7a] Load config and apply simulation patch...")
        config = load_config("config/station_2pumps.json")

        # Apply simulation patch
        sys.path.insert(0, os.path.dirname(__file__))
        from main import _patch_simulation
        _patch_simulation(config)

        from core.modbus_client import ModbusClient
        print("  PASS: Simulation patch applied")

        # 7b: Create modbus client and connect
        print("\n[7b] Simulated connect...")
        client = ModbusClient(config)
        result = client.connect()
        assert result is True
        assert client.connected is True
        print("  PASS: Simulated connection successful")

        # 7c: Receive simulated data
        print("\n[7c] Receive simulated data...")
        received_data = []
        client.data_received.connect(lambda readings: received_data.append(readings))
        client.start_polling()

        # Process events for a few poll cycles
        import time
        for _ in range(30):  # Wait up to 3 seconds
            app.processEvents()
            time.sleep(0.1)
            if len(received_data) >= 3:
                break

        assert len(received_data) >= 2, f"Expected >=2 data batches, got {len(received_data)}"
        first_batch = received_data[0]
        assert len(first_batch) == 5, f"Expected 5 sensors, got {len(first_batch)}"

        # Verify reading structure
        r = first_batch[0]
        assert hasattr(r, 'tag')
        assert hasattr(r, 'name')
        assert hasattr(r, 'category')
        assert hasattr(r, 'value')
        assert hasattr(r, 'unit')
        assert hasattr(r, 'timestamp')
        print(f"  PASS: Received {len(received_data)} batches, {len(first_batch)} sensors each")
        for reading in first_batch:
            print(f"    {reading.tag}: {reading.value} {reading.unit} ({reading.category})")

        # 7d: Simulated device control
        print("\n[7d] Simulated device control...")
        dev = config["devices"][0]
        result = client.set_device_on(dev, True)
        assert result is True
        result = client.set_device_setpoint(dev, 1500)
        assert result is True
        status = client.read_device_status(dev)
        assert status is not None  # Should return False in sim
        print(f"  PASS: Device control works (on=True, setpoint=1500, status={status})")

        # 7e: Disconnect
        print("\n[7e] Simulated disconnect...")
        client.disconnect()
        assert client.connected is False
        print("  PASS: Disconnected")

        # 7f: Full main window creation with simulation
        print("\n[7f] Create MainWindow in simulation mode...")
        from ui.main_window import MainWindow
        window = MainWindow(config=config, data_dir=tmp_dir)
        assert window is not None
        assert window._config == config
        print("  PASS: MainWindow created")

        # 7g: Simulate connect + data flow through main window
        print("\n[7g] Main window connect + data flow...")
        window._connect_action.trigger()

        for _ in range(30):
            app.processEvents()
            time.sleep(0.1)
            if window._trend_chart._tick_counter >= 2:
                break

        assert window._modbus.connected is True
        assert window._trend_chart._tick_counter >= 1, \
            f"Expected trend data, got tick={window._trend_chart._tick_counter}"
        print(f"  PASS: Data flowing, trend ticks = {window._trend_chart._tick_counter}")

        # 7h: Start recording through main window
        print("\n[7h] Start recording via main window...")
        window._recorder.start_recording()
        assert window._recorder.is_recording is True

        # Let a few data cycles run
        for _ in range(20):
            app.processEvents()
            time.sleep(0.1)

        info = window._recorder.get_recording_info()
        assert info is not None
        print(f"  PASS: Recording active, rows={info['rows']}")

        # 7i: Stop recording
        print("\n[7i] Stop recording...")
        final_path = window._recorder.stop_recording()
        assert final_path is not None
        assert os.path.isfile(final_path)
        print(f"  PASS: Recording saved to: {os.path.basename(final_path)}")

        # Verify saved data
        from openpyxl import load_workbook
        wb = load_workbook(final_path)
        for sheet in ["Temperature", "Pressure", "Flow"]:
            assert sheet in wb.sheetnames
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            print(f"    {sheet}: {len(rows)-1} data rows")
        wb.close()

        # Cleanup
        window._modbus.disconnect()
        window.close()

    except Exception as e:
        errors.append(f"7: {e}")
        import traceback
        traceback.print_exc()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    return errors


# ============================================================
# Main runner
# ============================================================
if __name__ == "__main__":
    os.chdir(os.path.dirname(os.path.abspath(__file__)))

    # Need virtual display for headless environment
    os.environ["QT_QPA_PLATFORM"] = "offscreen"

    all_errors = []
    test_funcs = [
        test_config_loader,
        test_data_recorder,
        test_modbus_client_logic,
        test_trend_chart,
        test_control_panel,
        test_history_viewer,
        test_simulation_integration,
    ]

    for test_func in test_funcs:
        try:
            errs = test_func()
            all_errors.extend(errs)
        except Exception as e:
            all_errors.append(f"{test_func.__name__}: CRASH: {e}")
            import traceback
            traceback.print_exc()

    # Summary
    print("\n" + "=" * 60)
    print("TEST SUMMARY")
    print("=" * 60)
    if all_errors:
        print(f"\nFAILED - {len(all_errors)} error(s):")
        for err in all_errors:
            print(f"  - {err}")
        sys.exit(1)
    else:
        print("\nALL TESTS PASSED!")
        sys.exit(0)
